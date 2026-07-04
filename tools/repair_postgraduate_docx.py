from __future__ import annotations

import os
import re
import shutil
import tempfile
import zipfile
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn as docx_qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "Agentic_SonarQube_Pipeline_Documentation_postgraduate_v12.docx"
OUTPUT = ROOT / "Agentic_SonarQube_Pipeline_Documentation_postgraduate_v13.docx"

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
XML_NS = "http://www.w3.org/XML/1998/namespace"
NS = {"w": W_NS, "r": R_NS}

FIGURE_TITLES = [
    "Αρχιτεκτονική Συστήματος Πολλαπλών Πρακτόρων",
    "Αντιστοίχιση Σχεδιαστικών Προτύπων Πρακτόρων στο Σύστημα",
    "Κλάση SonarCubeClient",
    "Μέθοδος get_issues()",
    "Παράδειγμα Δομής Δεδομένων JSON",
    "Σχεσιακό Μοντέλο Βάσης Δεδομένων (ER Diagram)",
    "Δομές μοντέλων δεδομένων",
    "Retry Loop και Reflection Mechanism",
    "Streamlit Dashboard: Κεντρική σελίδα με metric cards, cost summary και recent pipeline runs",
    "Σελίδα Run Pipeline: Επιλογή issues και εκκίνηση του Agent",
    "Σελίδα Issues: Λίστα ζητημάτων αντλημένη από τον SonarQube",
    "Analytics: Token usage, costs, success rates και cost breakdown ανά agent",
    "Docker Compose Τοπολογία",
    "Στιγμιότυπο εκτέλεσης του pipeline στο τερματικό",
    "Στιγμιότυπο εκτέλεσης - Εκκίνηση συνεδρίας και πρώτη απόπειρα διόρθωσης",
    "Στιγμιότυπο επιτυχούς επαλήθευσης και ελέγχου (Verification)",
    "Στιγμιότυπο εκτέλεσης - Δημιουργία και Εκτέλεση Δοκιμών (Tester Agent)",
    "Στιγμιότυπο εκτέλεσης - Διαδικασία Αξιολόγησης (Evaluator/Reviewer)",
    "Στιγμιότυπο εκτέλεσης - Πλήρης ροή επιτυχούς δοκιμής, αξιολόγησης και ενσωμάτωσης (Success Flow)",
]


def w(local: str) -> str:
    return f"{{{W_NS}}}{local}"


def r(local: str) -> str:
    return f"{{{R_NS}}}{local}"


def rel(local: str) -> str:
    return f"{{{REL_NS}}}{local}"


def text_of(el: etree._Element) -> str:
    return "".join(el.xpath(".//w:t/text()", namespaces=NS)).strip()


def style_id(p: etree._Element) -> str:
    style = p.find("./w:pPr/w:pStyle", namespaces=NS)
    return style.get(w("val")) if style is not None else ""


def delete_paragraph(paragraph: Paragraph) -> None:
    p = paragraph._element
    p.getparent().remove(p)
    paragraph._p = paragraph._element = None


def paragraph_after(ref: Paragraph, text: str = "", style_name: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    if style_name:
        p_pr = OxmlElement("w:pPr")
        p_style = OxmlElement("w:pStyle")
        p_style.set(docx_qn("w:val"), style_name.replace(" ", ""))
        p_pr.append(p_style)
        new_p.append(p_pr)
    if text:
        run = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = text
        run.append(t)
        new_p.append(run)
    ref._p.addnext(new_p)
    return Paragraph(new_p, ref._parent)


def add_image_to_paragraph(paragraph: Paragraph, image_path: Path, width: float = 6.0) -> Paragraph:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = None
    paragraph.paragraph_format.space_after = None
    paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    return paragraph


def find_paragraph(doc: Document, contains: str, style_prefix: str | None = None) -> Paragraph:
    for p in doc.paragraphs:
        if contains in (p.text or "") and (style_prefix is None or p.style.name.startswith(style_prefix)):
            return p
    raise RuntimeError(f"Could not find paragraph containing {contains!r}")


def extract_embedded_pipeline_screenshot(work_dir: Path) -> Path:
    out = work_dir / "terminal_pipeline_screenshot.png"
    with zipfile.ZipFile(INPUT) as zf:
        out.write_bytes(zf.read("word/media/image17.png"))
    return out


def stage_with_repaired_screenshots(stage_path: Path, work_dir: Path) -> None:
    doc = Document(str(INPUT))

    broken_start = None
    broken_end = None
    for idx, p in enumerate(doc.paragraphs):
        if "Screenshot Logs εκτέλεσης της ροής" in (p.text or ""):
            broken_start = idx
        if broken_start is not None and idx > broken_start and "Συμπεράσματα" in (p.text or "") and p.style.name.startswith("Heading"):
            broken_end = idx
            break

    if broken_start is not None and broken_end is not None:
        for p in list(doc.paragraphs)[broken_start:broken_end]:
            delete_paragraph(p)

    terminal_shot = extract_embedded_pipeline_screenshot(work_dir)

    # Restore the missing Analytics UI screenshot before Chapter 8.
    docker_heading = find_paragraph(doc, "8. Ανάπτυξη με Docker", "Heading")
    analytics_p = docker_heading.insert_paragraph_before("")
    add_image_to_paragraph(analytics_p, ROOT / "docs" / "screenshots" / "analytics.png", 6.5)

    # Put the terminal execution screenshot at the start of Chapter 10.
    chapter_10 = find_paragraph(doc, "10. Αποτελέσματα Εκτέλεσης", "Heading")
    term_p = paragraph_after(chapter_10)
    add_image_to_paragraph(term_p, terminal_shot, 6.0)

    # Put the first issue screenshot under 10.1, before the results table.
    results_heading = find_paragraph(doc, "10.1 Πίνακας Αποτελεσμάτων", "Heading")
    issue_p = paragraph_after(results_heading)
    add_image_to_paragraph(issue_p, ROOT / "first_issue_screenshot.png", 6.0)

    # Append QA/execution screenshots at the end of section 10.4.
    verification_heading = find_paragraph(doc, "10.5 Αποτελέσματα Επαλήθευσης", "Heading")
    for image_name, width in [
        ("final_run_screenshot.png", 6.0),
        ("tester_screenshot.png", 6.0),
        ("evaluator_screenshot.png", 6.0),
        ("success_flow_screenshot.png", 6.0),
    ]:
        p = verification_heading.insert_paragraph_before("")
        add_image_to_paragraph(p, ROOT / image_name, width)

    doc.save(str(stage_path))


def unzip_docx(docx_path: Path, out_dir: Path) -> None:
    if out_dir.exists():
        shutil.rmtree(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(docx_path, "r") as zf:
        zf.extractall(out_dir)


def zip_docx(in_dir: Path, out_docx_path: Path) -> None:
    if out_docx_path.exists():
        out_docx_path.unlink()
    with zipfile.ZipFile(out_docx_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for root, _dirs, files in os.walk(in_dir):
            for file_name in files:
                abs_path = Path(root) / file_name
                rel_path = abs_path.relative_to(in_dir)
                zf.write(abs_path, rel_path.as_posix())


def xml_space(t_el: etree._Element) -> None:
    if t_el.text and (t_el.text.startswith(" ") or t_el.text.endswith(" ")):
        t_el.set(f"{{{XML_NS}}}space", "preserve")


def run_text(text: str, bold: bool = False, italic: bool = False, size_half_points: int | None = None) -> etree._Element:
    run = etree.Element(w("r"))
    if bold or italic or size_half_points:
        r_pr = etree.SubElement(run, w("rPr"))
        r_fonts = etree.SubElement(r_pr, w("rFonts"))
        r_fonts.set(w("ascii"), "Times New Roman")
        r_fonts.set(w("hAnsi"), "Times New Roman")
        r_fonts.set(w("cs"), "Times New Roman")
        if bold:
            etree.SubElement(r_pr, w("b"))
        if italic:
            etree.SubElement(r_pr, w("i"))
        if size_half_points:
            sz = etree.SubElement(r_pr, w("sz"))
            sz.set(w("val"), str(size_half_points))
    t = etree.SubElement(run, w("t"))
    t.text = text
    xml_space(t)
    return run


def direct_title_paragraph(text: str) -> etree._Element:
    p = etree.Element(w("p"))
    p_pr = etree.SubElement(p, w("pPr"))
    spacing = etree.SubElement(p_pr, w("spacing"))
    spacing.set(w("before"), "360")
    spacing.set(w("after"), "180")
    p.append(run_text(text, bold=True, size_half_points=32))
    return p


def page_break_paragraph() -> etree._Element:
    p = etree.Element(w("p"))
    run = etree.SubElement(p, w("r"))
    br = etree.SubElement(run, w("br"))
    br.set(w("type"), "page")
    return p


def field_paragraph(instruction: str, placeholder: str = "Το πεδίο θα ενημερωθεί αυτόματα στο Word.") -> etree._Element:
    p = etree.Element(w("p"))
    r_begin = etree.SubElement(p, w("r"))
    fld_begin = etree.SubElement(r_begin, w("fldChar"))
    fld_begin.set(w("fldCharType"), "begin")
    fld_begin.set(w("dirty"), "true")

    r_instr = etree.SubElement(p, w("r"))
    instr = etree.SubElement(r_instr, w("instrText"))
    instr.set(f"{{{XML_NS}}}space", "preserve")
    instr.text = f" {instruction} "

    r_sep = etree.SubElement(p, w("r"))
    fld_sep = etree.SubElement(r_sep, w("fldChar"))
    fld_sep.set(w("fldCharType"), "separate")

    p.append(run_text(placeholder))

    r_end = etree.SubElement(p, w("r"))
    fld_end = etree.SubElement(r_end, w("fldChar"))
    fld_end.set(w("fldCharType"), "end")
    return p


def caption_paragraph(label: str, seq_name: str, title: str) -> etree._Element:
    p = etree.Element(w("p"))
    p_pr = etree.SubElement(p, w("pPr"))
    p_style = etree.SubElement(p_pr, w("pStyle"))
    p_style.set(w("val"), "Caption")
    jc = etree.SubElement(p_pr, w("jc"))
    jc.set(w("val"), "center")

    p.append(run_text(f"{label} ", italic=True))

    r_begin = etree.SubElement(p, w("r"))
    fld_begin = etree.SubElement(r_begin, w("fldChar"))
    fld_begin.set(w("fldCharType"), "begin")

    r_instr = etree.SubElement(p, w("r"))
    instr = etree.SubElement(r_instr, w("instrText"))
    instr.set(f"{{{XML_NS}}}space", "preserve")
    instr.text = f" SEQ {seq_name} \\* ARABIC "

    r_sep = etree.SubElement(p, w("r"))
    fld_sep = etree.SubElement(r_sep, w("fldChar"))
    fld_sep.set(w("fldCharType"), "separate")

    p.append(run_text("0", italic=True))

    r_end = etree.SubElement(p, w("r"))
    fld_end = etree.SubElement(r_end, w("fldChar"))
    fld_end.set(w("fldCharType"), "end")

    p.append(run_text(f": {title}", italic=True))
    return p


def replace_front_matter(body: etree._Element) -> None:
    children = list(body)
    start = None
    end = None
    for idx, child in enumerate(children):
        if child.tag == w("p") and text_of(child) == "Πίνακας Περιεχομένων":
            start = idx
        if start is not None and idx > start and child.tag == w("p") and text_of(child) == "Περίληψη":
            end = idx
            break
    if start is None or end is None:
        raise RuntimeError("Could not locate front matter block")

    for child in children[start:end]:
        body.remove(child)

    new_nodes = [
        direct_title_paragraph("Πίνακας Περιεχομένων"),
        field_paragraph('TOC \\o "1-3" \\h \\z \\u'),
        page_break_paragraph(),
        direct_title_paragraph("Ευρετήριο Εικόνων"),
        field_paragraph('TOC \\h \\z \\c "Figure"'),
        page_break_paragraph(),
        direct_title_paragraph("Ευρετήριο Πινάκων"),
        field_paragraph('TOC \\h \\z \\c "Table"'),
        page_break_paragraph(),
    ]
    for offset, node in enumerate(new_nodes):
        body.insert(start + offset, node)


def strip_heading_prefix(p: etree._Element, chars_to_remove: int) -> None:
    for t_el in p.xpath(".//w:t", namespaces=NS):
        if chars_to_remove <= 0:
            break
        txt = t_el.text or ""
        if chars_to_remove >= len(txt):
            t_el.text = ""
            chars_to_remove -= len(txt)
        else:
            t_el.text = txt[chars_to_remove:]
            chars_to_remove = 0
        xml_space(t_el)


def ensure_ppr(p: etree._Element) -> etree._Element:
    p_pr = p.find("./w:pPr", namespaces=NS)
    if p_pr is None:
        p_pr = etree.Element(w("pPr"))
        p.insert(0, p_pr)
    return p_pr


def add_heading_numbering(p: etree._Element, num_id: int, ilvl: int) -> None:
    p_pr = ensure_ppr(p)
    for existing in p_pr.findall("w:numPr", namespaces=NS):
        p_pr.remove(existing)
    num_pr = etree.Element(w("numPr"))
    ilvl_el = etree.SubElement(num_pr, w("ilvl"))
    ilvl_el.set(w("val"), str(ilvl))
    num_id_el = etree.SubElement(num_pr, w("numId"))
    num_id_el.set(w("val"), str(num_id))
    p_style = p_pr.find("w:pStyle", namespaces=NS)
    if p_style is not None:
        p_pr.insert(p_pr.index(p_style) + 1, num_pr)
    else:
        p_pr.insert(0, num_pr)


def install_heading_numbering(unzipped: Path) -> int:
    numbering_path = unzipped / "word" / "numbering.xml"
    parser = etree.XMLParser(remove_blank_text=False)
    if numbering_path.exists():
        tree = etree.parse(str(numbering_path), parser)
        root = tree.getroot()
    else:
        root = etree.Element(w("numbering"), nsmap={"w": W_NS})
        tree = etree.ElementTree(root)

    abstract_ids = [int(el.get(w("abstractNumId"))) for el in root.findall("w:abstractNum", namespaces=NS) if (el.get(w("abstractNumId")) or "").isdigit()]
    num_ids = [int(el.get(w("numId"))) for el in root.findall("w:num", namespaces=NS) if (el.get(w("numId")) or "").isdigit()]
    abstract_id = (max(abstract_ids) + 1) if abstract_ids else 1
    num_id = (max(num_ids) + 1) if num_ids else 1

    abstract = etree.Element(w("abstractNum"))
    abstract.set(w("abstractNumId"), str(abstract_id))
    multi = etree.SubElement(abstract, w("multiLevelType"))
    multi.set(w("val"), "hybridMultilevel")

    for level in range(4):
        lvl = etree.SubElement(abstract, w("lvl"))
        lvl.set(w("ilvl"), str(level))
        start = etree.SubElement(lvl, w("start"))
        start.set(w("val"), "1")
        fmt = etree.SubElement(lvl, w("numFmt"))
        fmt.set(w("val"), "decimal")
        if level > 0:
            restart = etree.SubElement(lvl, w("lvlRestart"))
            restart.set(w("val"), str(level))
        lvl_text = etree.SubElement(lvl, w("lvlText"))
        lvl_text.set(w("val"), ".".join(f"%{i}" for i in range(1, level + 2)) + ".")
        suff = etree.SubElement(lvl, w("suff"))
        suff.set(w("val"), "space")
        jc = etree.SubElement(lvl, w("lvlJc"))
        jc.set(w("val"), "left")

    root.append(abstract)
    num = etree.Element(w("num"))
    num.set(w("numId"), str(num_id))
    abstract_ref = etree.SubElement(num, w("abstractNumId"))
    abstract_ref.set(w("val"), str(abstract_id))
    root.append(num)

    tree.write(str(numbering_path), xml_declaration=True, encoding="UTF-8", standalone="yes")
    return num_id


def normalize_headings(root: etree._Element, num_id: int) -> None:
    pattern = re.compile(r"^\s*(\d+(?:\.\d+)*)(?:\.)?\s+")
    for p in root.findall(".//w:p", namespaces=NS):
        sid = style_id(p)
        if not sid.startswith("Heading"):
            continue
        text = text_of(p)
        match = pattern.match(text)
        if not match:
            continue
        parts = match.group(1).split(".")
        ilvl = min(len(parts) - 1, 3)
        strip_heading_prefix(p, len(match.group(0)))
        add_heading_numbering(p, num_id, ilvl)


def transform_table_captions(root: etree._Element) -> None:
    caption_re = re.compile(r"^Πίνακας\s+\d+\s*[:—-]\s*(.+)$")
    for p in list(root.findall(".//w:p", namespaces=NS)):
        if style_id(p) != "Caption":
            continue
        text = text_of(p)
        match = caption_re.match(text)
        if not match:
            continue
        parent = p.getparent()
        idx = parent.index(p)
        parent.remove(p)
        parent.insert(idx, caption_paragraph("Πίνακας", "Table", match.group(1).strip()))


def rels_by_id(unzipped: Path) -> dict[str, str]:
    rels_path = unzipped / "word" / "_rels" / "document.xml.rels"
    tree = etree.parse(str(rels_path))
    out = {}
    for rel_el in tree.findall(f".//{{{REL_NS}}}Relationship"):
        out[rel_el.get("Id")] = rel_el.get("Target", "")
    return out


def add_figure_captions(root: etree._Element, unzipped: Path) -> None:
    rels = rels_by_id(unzipped)
    by_paragraph: dict[etree._Element, list[str]] = defaultdict(list)

    for p in root.findall(".//w:p", namespaces=NS):
        embed_ids = p.xpath(".//@r:embed", namespaces=NS)
        for embed_id in embed_ids:
            target = rels.get(embed_id, "")
            if not target or target.endswith("image1.jpeg"):
                continue
            if not target.startswith("media/"):
                continue
            by_paragraph[p].append(target)

    ordered_titles = iter(FIGURE_TITLES)
    used = 0
    for p in list(root.findall(".//w:p", namespaces=NS)):
        count = len(by_paragraph.get(p, []))
        if not count:
            continue
        titles = []
        for _ in range(count):
            try:
                titles.append(next(ordered_titles))
            except StopIteration as exc:
                raise RuntimeError("More figure images than available titles") from exc
        used += len(titles)
        for title in reversed(titles):
            p.addnext(caption_paragraph("Εικόνα", "Figure", title))

    if used != len(FIGURE_TITLES):
        raise RuntimeError(f"Expected {len(FIGURE_TITLES)} figure captions, inserted {used}")


def set_update_fields_on_open(unzipped: Path) -> None:
    settings_path = unzipped / "word" / "settings.xml"
    parser = etree.XMLParser(remove_blank_text=False)
    if settings_path.exists():
        tree = etree.parse(str(settings_path), parser)
        root = tree.getroot()
    else:
        root = etree.Element(w("settings"), nsmap={"w": W_NS})
        tree = etree.ElementTree(root)
    existing = root.find("w:updateFields", namespaces=NS)
    if existing is None:
        existing = etree.Element(w("updateFields"))
        root.insert(0, existing)
    existing.set(w("val"), "true")
    tree.write(str(settings_path), xml_declaration=True, encoding="UTF-8", standalone="yes")


def patch_ooxml(stage_path: Path, out_path: Path, work_dir: Path) -> None:
    unzipped = work_dir / "unzipped"
    unzip_docx(stage_path, unzipped)

    num_id = install_heading_numbering(unzipped)

    doc_xml = unzipped / "word" / "document.xml"
    parser = etree.XMLParser(remove_blank_text=False)
    tree = etree.parse(str(doc_xml), parser)
    root = tree.getroot()
    body = root.find("w:body", namespaces=NS)
    if body is None:
        raise RuntimeError("Document body not found")

    replace_front_matter(body)
    normalize_headings(root, num_id)
    transform_table_captions(root)
    add_figure_captions(root, unzipped)

    tree.write(str(doc_xml), xml_declaration=True, encoding="UTF-8", standalone="yes")
    set_update_fields_on_open(unzipped)
    zip_docx(unzipped, out_path)


def main() -> None:
    if not INPUT.exists():
        raise FileNotFoundError(INPUT)

    with tempfile.TemporaryDirectory(prefix="docx_repair_", dir=str(ROOT)) as tmp:
        work_dir = Path(tmp)
        stage_path = work_dir / "stage.docx"
        stage_with_repaired_screenshots(stage_path, work_dir)
        patch_ooxml(stage_path, OUTPUT, work_dir)

    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
